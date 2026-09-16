"""Three-sheet Excel export for the funnel report: Funnel, Segments, Events."""

from __future__ import annotations

import json
import os
from datetime import datetime

from openpyxl import Workbook

import config
from core.xlsx_helpers import BOLD, CENTER, LEFT, NORMAL, RIGHT, write_header
from db import events_since
from features.funnel.analyzer import DIMENSIONS, STAGES, FunnelReport


def export(report: FunnelReport, days: int) -> str:
    os.makedirs(config.EXPORTS_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y-%m-%d_%H%M")
    path = os.path.join(config.EXPORTS_DIR, f"funnel_{ts}.xlsx")

    wb = Workbook()
    wb.remove(wb.active)
    _xl_funnel(wb, report, days)
    _xl_segments(wb, report)
    _xl_events(wb, report)
    wb.save(path)
    wb.save(os.path.join(config.EXPORTS_DIR, "funnel_latest.xlsx"))
    return path


def _xl_funnel(wb: Workbook, report: FunnelReport, days: int) -> None:
    ws = wb.create_sheet("Funnel")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14
    rows = [
        ("Window", f"Last {days} days (since {report.since[:10]})"),
        ("Jobs / events", f"{report.jobs} / {report.events}"),
        ("Win rate (shrunk)", round(report.win.shrunk * 100, 1)),
        ("Win rate (raw)", round(report.win.raw * 100, 1) if report.win.raw is not None else None),
        ("Submitted / hired", f"{report.win.submitted} / {report.win.hired}"),
        ("Insufficient sample (<5 submitted)", "yes" if report.win.insufficient else "no"),
        ("Connects spent", report.connects_spent),
        ("Connect price (USD)", config.CONNECT_PRICE_USD),
        ("Cost per hire (USD)", report.cost_per_hire_usd),
        ("Median submitted bid", report.med_submitted_bid),
        ("Median winning bid", report.med_winning_bid),
        ("Last ingest", report.last_ingested_at or "never"),
    ]
    for i, (k, v) in enumerate(rows, 1):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v).font = NORMAL
    start = len(rows) + 2
    ws.cell(row=start, column=1, value="Stage").font = BOLD
    ws.cell(row=start, column=2, value="Jobs").font = BOLD
    ws.cell(row=start, column=3, value="Conversion").font = BOLD
    prev = None
    for i, stage in enumerate(STAGES, start + 1):
        n = report.stages[stage]
        ws.cell(row=i, column=1, value=stage.title())
        ws.cell(row=i, column=2, value=n).alignment = RIGHT
        ws.cell(row=i, column=3, value=round(n / prev, 3) if prev else None).alignment = RIGHT
        prev = n
    for j, (stage, n) in enumerate(report.side_exits.items(), start + len(STAGES) + 1):
        ws.cell(row=j, column=1, value=stage.title())
        ws.cell(row=j, column=2, value=n).alignment = RIGHT


def _xl_segments(wb: Workbook, report: FunnelReport) -> None:
    ws = wb.create_sheet("Segments")
    headers = ["Dimension", "Value", "Notified", "Drafted", "Submitted", "Hired",
               "Win rate (shrunk) %", "Win rate (raw) %", "Insufficient"]  # fmt: skip
    write_header(ws, headers, [16, 28, 10, 10, 11, 8, 18, 16, 12])
    r = 2
    for dim in DIMENSIONS:
        for row in report.segments.get(dim, []):
            values = [
                dim, row.value, row.notified, row.drafted, row.submitted, row.hired,
                round(row.win.shrunk * 100, 1),
                round(row.win.raw * 100, 1) if row.win.raw is not None else None,
                "yes" if row.win.insufficient else "no",
            ]  # fmt: skip
            for c, v in enumerate(values, 1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.alignment = LEFT if c <= 2 else RIGHT
            r += 1
    ws.freeze_panes = "A2"


def _xl_events(wb: Workbook, report: FunnelReport) -> None:
    ws = wb.create_sheet("Events")
    headers = [
        "ts",
        "event",
        "job_id",
        "source",
        "bid",
        "bid type",
        "connects",
        "title",
        "matched",
        "note",
    ]
    write_header(ws, headers, [20, 11, 22, 10, 8, 9, 9, 50, 22, 30])
    for i, e in enumerate(events_since(report.since), 2):
        try:
            meta = json.loads(e["meta_json"] or "{}")
        except ValueError:
            meta = {}
        values = [
            e["ts"], e["event"], e["job_id"], e["source"], e["bid_amount"], e["bid_type"],
            e["connects_spent"], meta.get("title", ""), meta.get("matched", ""), meta.get("note", ""),
        ]  # fmt: skip
        for c, v in enumerate(values, 1):
            ws.cell(row=i, column=c, value=v).alignment = CENTER if c in (2, 4) else LEFT
    ws.freeze_panes = "A2"
