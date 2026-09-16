"""6-sheet Excel export for the n8n report."""

from __future__ import annotations

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

import config
from core.models import Job, N8nReport
from core.rich_helpers import budget_label
from core.scoring import get_scoring
from core.xlsx_helpers import BOLD, CENTER, HDR_FILL, HDR_FONT, LEFT, NORMAL, RIGHT, write_header


def export(report: N8nReport, days: int, *, data_as_of: str = "") -> str:
    """Build workbook; return path. Also writes `n8n_demand_latest.xlsx`."""
    os.makedirs(config.EXPORTS_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y-%m-%d_%H%M")
    path = os.path.join(config.EXPORTS_DIR, f"n8n_demand_{ts}.xlsx")

    wb = Workbook()
    wb.remove(wb.active)

    _xl_summary(wb, report, days, data_as_of)
    _xl_workflow(wb, report)
    _xl_industry(wb, report)
    _xl_stack(wb, report)
    _xl_matrix(wb, report)
    _xl_jobs(wb, report)

    wb.save(path)
    wb.save(os.path.join(config.EXPORTS_DIR, "n8n_demand_latest.xlsx"))
    return path


def _xl_summary(wb: Workbook, report: N8nReport, days: int, data_as_of: str = "") -> None:
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 50
    rows = [
        ("Data as of", data_as_of or "unknown"),
        ("Window", f"Last {days} days"),
        ("Total n8n jobs", report.total_jobs),
        ("Top workflow", report.workflow_count[0][0] if report.workflow_count else "—"),
        ("Top industry", report.industry_count[0][0] if report.industry_count else "—"),
        ("Top stack", report.stack_count[0][0] if report.stack_count else "—"),
        ("Unclassified industry", report.unclassified_industry),
        ("Unclassified workflow", report.unclassified_workflow),
        *get_scoring().summary_rows(),
    ]
    for i, (k, v) in enumerate(rows, 1):
        ws.cell(row=i, column=1, value=k).font = BOLD
        ws.cell(row=i, column=2, value=v).font = NORMAL


def _xl_workflow(wb: Workbook, report: N8nReport) -> None:
    ws = wb.create_sheet("By Workflow")
    headers = [
        "#",
        "Workflow",
        "Jobs",
        "Share %",
        "Med Hourly $",
        "Hourly p25",
        "Hourly p75",
        "n Hourly",
        "Med Fixed $",
        "Fixed p25",
        "Fixed p75",
        "n Fixed",
        "Med Proposals",
        "Verified %",
    ]
    widths = [4, 28, 7, 9, 13, 11, 11, 9, 13, 11, 11, 8, 13, 11]
    write_header(ws, headers, widths)

    total = report.total_jobs
    for i, (name, count) in enumerate(report.workflow_count, 1):
        s = report.workflow_stats[name]
        r = i + 1
        ws.cell(row=r, column=1, value=i).alignment = CENTER
        ws.cell(row=r, column=2, value=name).font = BOLD
        values = [
            count,
            round(count / total * 100, 1) if total else 0,
            round(s.med_hourly, 2) if s.med_hourly else None,
            round(s.hourly_p25, 2) if s.hourly_p25 else None,
            round(s.hourly_p75, 2) if s.hourly_p75 else None,
            s.n_hourly,
            round(s.med_fixed, 2) if s.med_fixed else None,
            round(s.fixed_p25, 2) if s.fixed_p25 else None,
            round(s.fixed_p75, 2) if s.fixed_p75 else None,
            s.n_fixed,
            round(s.med_proposals, 1) if s.med_proposals else None,
            s.verified_pct,
        ]
        for c, val in enumerate(values, 3):
            ws.cell(row=r, column=c, value=val).alignment = RIGHT
    last = len(report.workflow_count) + 1
    if last > 1:
        ws.conditional_formatting.add(
            f"C2:C{last}",
            ColorScaleRule(
                start_type="min", start_color="FFFFFF", end_type="max", end_color="4472C4"
            ),
        )
    ws.freeze_panes = "A2"


def _xl_count_sheet(
    wb: Workbook,
    name: str,
    label_col: str,
    items: list[tuple[str, int]],
    total: int,
) -> None:
    ws = wb.create_sheet(name)
    write_header(ws, ["#", label_col, "Jobs", "Share %"], [4, 28, 8, 10])
    for i, (label, count) in enumerate(items, 1):
        r = i + 1
        ws.cell(row=r, column=1, value=i).alignment = CENTER
        ws.cell(row=r, column=2, value=label).font = BOLD
        ws.cell(row=r, column=3, value=count).alignment = RIGHT
        ws.cell(
            row=r, column=4, value=round(count / total * 100, 1) if total else 0
        ).alignment = RIGHT
    ws.freeze_panes = "A2"


def _xl_industry(wb: Workbook, report: N8nReport) -> None:
    _xl_count_sheet(wb, "By Industry", "Industry", report.industry_count, report.total_jobs)


def _xl_stack(wb: Workbook, report: N8nReport) -> None:
    _xl_count_sheet(wb, "By Stack", "Tool / Stack", report.stack_count, report.total_jobs)


def _xl_matrix(wb: Workbook, report: N8nReport) -> None:
    ws = wb.create_sheet("Industry x Workflow")
    industries = [i for i, _ in report.industry_count]
    workflows = [w for w, _ in report.workflow_count]
    industries = [i for i in industries if any(report.matrix.get(i, {}).values())]
    workflows = [
        w for w in workflows if any(report.matrix.get(i, {}).get(w, 0) for i in industries)
    ]
    if not industries or not workflows:
        return

    cell = ws.cell(row=1, column=1, value="Industry \\ Workflow")
    cell.fill = HDR_FILL
    cell.font = HDR_FONT
    ws.column_dimensions["A"].width = 26
    for c, wf in enumerate(workflows, 2):
        cell = ws.cell(row=1, column=c, value=wf)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        ws.column_dimensions[get_column_letter(c)].width = 18
    for r, ind in enumerate(industries, 2):
        ws.cell(row=r, column=1, value=ind).font = BOLD
        for c, wf in enumerate(workflows, 2):
            ws.cell(row=r, column=c, value=report.matrix.get(ind, {}).get(wf, 0)).alignment = CENTER
    ws.freeze_panes = "B2"
    last_col = get_column_letter(len(workflows) + 1)
    ws.conditional_formatting.add(
        f"B2:{last_col}{len(industries) + 1}",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color="FFFFFF",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFD966",
            end_type="max",
            end_color="C00000",
        ),
    )


def _xl_jobs(wb: Workbook, report: N8nReport) -> None:
    ws = wb.create_sheet("Jobs")
    headers = [
        "Score",
        "Posted",
        "Title",
        "Industries",
        "Workflows",
        "Stacks",
        "Budget",
        "Proposals",
        "Country",
        "Verified",
        "Hires",
        "URL",
        "Job ID",
    ]
    widths = [6, 12, 60, 22, 28, 28, 14, 10, 14, 9, 7, 36, 26]
    write_header(ws, headers, widths)

    for i, j in enumerate(report.jobs, 1):
        r = i + 1
        ws.cell(row=r, column=1, value=round(j.opp_score)).alignment = CENTER
        ws.cell(row=r, column=2, value=(j.published_at or "")[:10]).alignment = CENTER
        ws.cell(row=r, column=3, value=j.display_title).alignment = LEFT
        ws.cell(row=r, column=4, value=", ".join(j.industries)).alignment = LEFT
        ws.cell(row=r, column=5, value=", ".join(j.workflows)).alignment = LEFT
        ws.cell(row=r, column=6, value=", ".join(j.stacks)).alignment = LEFT
        ws.cell(row=r, column=7, value=_budget_label(j)).alignment = RIGHT
        ws.cell(row=r, column=8, value=j.total_applicants).alignment = RIGHT
        ws.cell(row=r, column=9, value=j.client_country).alignment = LEFT
        ws.cell(row=r, column=10, value="Yes" if j.client_verified else "No").alignment = CENTER
        ws.cell(row=r, column=11, value=j.client_total_hires).alignment = RIGHT
        url_cell = ws.cell(row=r, column=12, value=j.url)
        url_cell.hyperlink = j.url
        url_cell.alignment = LEFT
        ws.cell(row=r, column=13, value=j.id).alignment = LEFT
    ws.freeze_panes = "A2"
    if report.jobs:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(report.jobs) + 1}"
        ws.conditional_formatting.add(
            f"A2:A{len(report.jobs) + 1}",
            ColorScaleRule(
                start_type="min",
                start_color="FFFFFF",
                mid_type="percentile",
                mid_value=50,
                mid_color="FFE699",
                end_type="max",
                end_color="70AD47",
            ),
        )


def _budget_label(j: Job) -> str:
    return budget_label(j)
