"""Two-sheet Excel export: Bands, Your bids."""

from __future__ import annotations

import os
from datetime import datetime

from openpyxl import Workbook

import config
from core.xlsx_helpers import LEFT, RIGHT, write_header
from features.bands.analyzer import ALL, BandsReport


def export(report: BandsReport, days: int) -> str:
    os.makedirs(config.EXPORTS_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y-%m-%d_%H%M")
    path = os.path.join(config.EXPORTS_DIR, f"bands_{ts}.xlsx")
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Bands")
    write_header(ws, ["Level", "Workflow", "Segment", "Experience", "Type", "n", "p25", "Median", "p75"],
                 [12, 28, 12, 14, 8, 6, 10, 10, 10])  # fmt: skip
    for i, r in enumerate(report.rows, 2):
        values = [r.level, "ALL" if r.workflow == ALL else r.workflow, "" if r.client_segment == ALL else r.client_segment,
                  "" if r.experience == ALL else r.experience, r.budget_type, r.n, r.p25, r.p50, r.p75]  # fmt: skip
        for c, v in enumerate(values, 1):
            ws.cell(row=i, column=c, value=v).alignment = LEFT if c <= 5 else RIGHT
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Your bids")
    write_header(ws, ["When", "Job", "Workflow", "Segment", "Type", "Bid", "Band p25", "Band median", "Band p75", "Band level", "Position", "Outcome"],
                 [12, 22, 24, 12, 8, 10, 10, 12, 10, 12, 10, 10])  # fmt: skip
    for i, b in enumerate(report.bids, 2):
        values = [b.ts[:10], b.job_id, b.workflow, b.client_segment, b.bid_type, b.bid_amount,
                  b.band_p25, b.band_p50, b.band_p75, b.band_level, b.position, b.outcome]  # fmt: skip
        for c, v in enumerate(values, 1):
            ws.cell(row=i, column=c, value=v).alignment = LEFT if c <= 5 or c >= 10 else RIGHT
    ws.freeze_panes = "A2"

    wb.save(path)
    wb.save(os.path.join(config.EXPORTS_DIR, "bands_latest.xlsx"))
    return path
