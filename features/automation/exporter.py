"""Excel export: the six n8n sheets plus By Platform and Platform x Workflow."""

from __future__ import annotations

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

import config
from core.xlsx_helpers import BOLD, CENTER, HDR_FILL, HDR_FONT
from features.automation.analyzer import AutomationReport
from features.n8n import exporter as n8n


def export(report: AutomationReport, days: int, *, data_as_of: str = "") -> str:
    os.makedirs(config.EXPORTS_DIR, exist_ok=True)
    ts = datetime.now().strftime("%Y-%m-%d_%H%M")
    path = os.path.join(config.EXPORTS_DIR, f"automation_{ts}.xlsx")

    wb = Workbook()
    wb.remove(wb.active)
    n8n._xl_summary(wb, report.base, days, data_as_of)
    ws = wb["Summary"]
    r = ws.max_row + 1
    ws.cell(row=r, column=1, value="Platforms").font = BOLD
    ws.cell(row=r, column=2, value=", ".join(report.platforms) or "all")
    n8n._xl_count_sheet(wb, "By Platform", "Platform", report.platform_count, report.total_jobs)
    n8n._xl_workflow(wb, report.base)
    n8n._xl_industry(wb, report.base)
    n8n._xl_stack(wb, report.base)
    _xl_platform_matrix(wb, report)
    n8n._xl_matrix(wb, report.base)
    n8n._xl_jobs(wb, report.base)

    wb.save(path)
    wb.save(os.path.join(config.EXPORTS_DIR, "automation_latest.xlsx"))
    return path


def _xl_platform_matrix(wb: Workbook, report: AutomationReport) -> None:
    ws = wb.create_sheet("Platform x Workflow")
    platforms = [
        p for p, _ in report.platform_count if any(report.platform_matrix.get(p, {}).values())
    ]
    workflows = [w for w, _ in report.base.workflow_count]
    workflows = [
        w for w in workflows if any(report.platform_matrix.get(p, {}).get(w, 0) for p in platforms)
    ]
    if not platforms or not workflows:
        return
    cell = ws.cell(row=1, column=1, value="Platform \\ Workflow")
    cell.fill, cell.font = HDR_FILL, HDR_FONT
    ws.column_dimensions["A"].width = 24
    for c, wf in enumerate(workflows, 2):
        h = ws.cell(row=1, column=c, value=wf)
        h.fill, h.font, h.alignment = HDR_FILL, HDR_FONT, CENTER
        ws.column_dimensions[get_column_letter(c)].width = 18
    for r, p in enumerate(platforms, 2):
        ws.cell(row=r, column=1, value=p).font = BOLD
        for c, wf in enumerate(workflows, 2):
            ws.cell(
                row=r, column=c, value=report.platform_matrix.get(p, {}).get(wf, 0)
            ).alignment = CENTER
    ws.freeze_panes = "B2"
    last_col = get_column_letter(len(workflows) + 1)
    ws.conditional_formatting.add(
        f"B2:{last_col}{len(platforms) + 1}",
        ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF", mid_type="percentile",
                       mid_value=50, mid_color="FFD966", end_type="max", end_color="C00000"),
    )  # fmt: skip
