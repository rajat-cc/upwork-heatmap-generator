"""openpyxl primitives shared across feature exporters."""
from __future__ import annotations

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# ── Palette ─────────────────────────────────────────────────────────────────
HDR_FILL = PatternFill("solid", fgColor="1F3864")        # dark navy
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
ALT_FILL = PatternFill("solid", fgColor="EEF2F7")        # light blue-grey
GREEN_FILL = PatternFill("solid", fgColor="D6EAD6")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
RED_FILL = PatternFill("solid", fgColor="F4CCCC")

BOLD = Font(bold=True, size=10)
NORMAL = Font(size=10)
ITALIC = Font(italic=True, size=10)

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right", vertical="center")


def write_header(ws: Worksheet, headers: list[str], widths: list[int]) -> None:
    """Write a styled header row + set column widths."""
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        ws.column_dimensions[get_column_letter(c)].width = w


def style_title_row(ws: Worksheet, row: int, ncols: int, title: str, *, insert: bool = True) -> None:
    """Insert a bold title row above the header at `row`."""
    if insert:
        ws.insert_rows(row)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = Font(bold=True, size=12, color="1F3864")
    cell.alignment = LEFT
    ws.row_dimensions[row].height = 20
